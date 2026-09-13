"""Standalone PDF -> DOCX converter.

This module deliberately does not contain a second PDF parser.  It calls the
same _imp_convert_pdf() path used by the SDYnotes importer, then turns that
parser's page elements into positioned WordprocessingML.  The converter has its
own temporary upload/job/output directories and never creates a note or a user
session.
"""
from copy import deepcopy
from html.parser import HTMLParser
import html
import json
import os
import re
import threading
import time
import uuid

from flask import jsonify, request, send_file

from .common import BASE_DIR, IMG_DIR
from .core import app
from . import importer as _engine

try:
    from .importer import _imp_uid
except Exception:  # pragma: no cover - importer always provides this in production
    def _imp_uid(prefix):
        return f"{prefix}_{uuid.uuid4().hex[:9]}"

CONVERTER_UPLOAD_DIR = os.path.join(BASE_DIR, "converter_uploads")
CONVERTER_JOB_DIR = os.path.join(BASE_DIR, "converter_jobs")
CONVERTER_OUT_DIR = os.path.join(BASE_DIR, "converter_output")
CONVERTER_TMP_DIR = os.path.join(BASE_DIR, "converter_tmp")
CONVERTER_MAX_MB = int(os.environ.get("SDY_CONVERTER_MAX_MB", "120") or 120)
CONVERTER_JOB_TTL = 3600
_CONVERTER_ID = re.compile(r"^[0-9a-f]{12,32}$")

for _d in (CONVERTER_UPLOAD_DIR, CONVERTER_JOB_DIR, CONVERTER_OUT_DIR,
           CONVERTER_TMP_DIR):
    os.makedirs(_d, exist_ok=True)


def _safe_id(value):
    value = str(value or "")
    return value if _CONVERTER_ID.fullmatch(value) else ""


def _job_path(jid):
    return os.path.join(CONVERTER_JOB_DIR, f"{jid}.json")


def _output_path(jid):
    return os.path.join(CONVERTER_OUT_DIR, f"{jid}.docx")


def _write_job(jid, **values):
    path = _job_path(jid)
    current = {}
    try:
        if os.path.exists(path):
            with open(path, encoding="utf-8") as fp:
                current = json.load(fp)
    except (OSError, ValueError):
        current = {}
    current.update(values)
    current["updated"] = time.time()
    tmp = f"{path}.tmp.{uuid.uuid4().hex[:6]}"
    with open(tmp, "w", encoding="utf-8") as fp:
        json.dump(current, fp, ensure_ascii=False)
    os.replace(tmp, path)


def _read_job(jid):
    try:
        with open(_job_path(jid), encoding="utf-8") as fp:
            return json.load(fp)
    except (OSError, ValueError):
        return None


def _sweep_converter_jobs():
    now = time.time()
    for directory in (CONVERTER_UPLOAD_DIR, CONVERTER_JOB_DIR,
                      CONVERTER_OUT_DIR, CONVERTER_TMP_DIR):
        try:
            for name in os.listdir(directory):
                path = os.path.join(directory, name)
                if now - os.path.getmtime(path) > CONVERTER_JOB_TTL:
                    if os.path.isfile(path):
                        os.remove(path)
        except (OSError, ValueError):
            pass


def _asset_name(url):
    match = re.fullmatch(r"/api/import/img/([0-9a-f]{8,32}\.(?:png|jpg|jpeg|svg))",
                         str(url or ""))
    return match.group(1) if match else None


def _asset_path(url, jid):
    """Resolve only parser-created asset names; never accept an arbitrary path."""
    name = _asset_name(url)
    if not name:
        return None
    path = os.path.join(IMG_DIR, name)
    if not os.path.isfile(path):
        return None
    if not name.endswith(".svg"):
        return path
    # Word versions in the wild do not consistently display SVG in DOCX.  The
    # parser's SVG is rendered with MuPDF (the same renderer used by the import
    # engine), not with a second PDF/image interpretation.
    out = os.path.join(CONVERTER_TMP_DIR, f"{jid}_{name[:-4]}.png")
    if os.path.exists(out):
        return out
    try:
        import pymupdf
        with open(path, "rb") as fp:
            raw = fp.read()
        with pymupdf.open(stream=raw, filetype="svg") as svg_doc:
            svg_doc[0].get_pixmap(alpha=False, dpi=180).save(out)
        return out if os.path.isfile(out) else None
    except Exception:
        return None


class _HtmlRuns(HTMLParser):
    """Small, purpose-built parser for importer text element HTML."""
    def __init__(self, default_font="Arial", default_px=12):
        super().__init__(convert_charrefs=True)
        self.default = {
            "font": _word_font(default_font or "Arial"),
            "px": float(default_px or 12),
            "bold": False,
            "italic": False,
            "underline": False,
            "color": None,
        }
        self.stack = [dict(self.default)]
        self.runs = []
        # PDF text elements contain absolutely positioned spans.  Word text
        # boxes are also positioned, but their internal runs are normal flow;
        # restore line breaks from the recorded baseline instead of collapsing
        # a multi-line paragraph into one line.
        self._last_base = None
        self._last_px = float(default_px or 12)
        self._pending_prefix = ""

    @staticmethod
    def _style(attrs):
        style = dict(attrs).get("style", "") or ""
        out = {}
        for key, value in re.findall(r"([\w-]+)\s*:\s*([^;]+)", style):
            out[key.lower()] = value.strip()
        return out

    @staticmethod
    def _font(value):
        return _word_font(value)

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        state = dict(self.stack[-1])
        attrs_d = dict(attrs)
        if tag == "span":
            styles = self._style(attrs)
            state["font"] = self._font(styles.get("font-family", state["font"]))
            raw_px = attrs_d.get("data-fs") or styles.get("font-size")
            if raw_px:
                try:
                    state["px"] = float(re.sub(r"[^0-9.+-]", "", raw_px))
                except ValueError:
                    pass
            if styles.get("font-weight", "").lower() in ("bold", "700", "800", "900"):
                state["bold"] = True
            if styles.get("font-style", "").lower() == "italic":
                state["italic"] = True
            color = styles.get("color", "")
            match = re.search(r"#([0-9a-f]{6})", color, re.I)
            if match:
                state["color"] = match.group(1)
            classes = (attrs_d.get("class") or "").split()
            if "zsp" in classes:
                state["_space"] = True
            # data-pdf-base is the baseline relative to this text element. A
            # small baseline change is a superscript/subscript; a large one is
            # a new PDF line. This deliberately avoids using top, which would
            # mistake a legitimate superscript for a paragraph break.
            if "zsp" not in classes and attrs_d.get("data-pdf-base") is not None:
                try:
                    base = float(attrs_d["data-pdf-base"])
                    if self._last_base is not None:
                        threshold = max(4.0, max(self._last_px, state["px"]) * .75)
                        if abs(base - self._last_base) > threshold:
                            self._pending_prefix += "\n"
                    self._last_base = base
                    self._last_px = state["px"]
                except (TypeError, ValueError):
                    pass
        elif tag in ("b", "strong"):
            state["bold"] = True
        elif tag in ("i", "em"):
            state["italic"] = True
        elif tag == "u":
            state["underline"] = True
        self.stack.append(state)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if len(self.stack) > 1:
            self.stack.pop()

    def handle_data(self, data):
        if not data:
            return
        state = dict(self.stack[-1])
        if state.pop("_space", False) and not data.strip():
            data = " "
        if self._pending_prefix:
            data = self._pending_prefix + data
            self._pending_prefix = ""
        self.runs.append((data, state))

    def handle_entityref(self, name):
        self.handle_data(html.unescape(f"&{name};"))



def _word_font(value):
    """Translate importer/CSS font names into fonts Word is likely to have.

    The SDYnotes renderer can use bundled web fonts such as ``SDY Computer
    Modern``.  DOCX run properties, however, name installed Office fonts; if we
    write the site-only family literally, Word silently substitutes and the
    geometry drifts.  Prefer common Word fonts from a CSS fallback list and map
    importer font IDs directly.
    """
    raw = str(value or "").strip()
    key = re.sub(r"[^a-z0-9]", "", raw.lower())
    id_map = {
        "times": "Times New Roman",
        "cmroman": "Times New Roman",
        "arial": "Arial",
        "mono": "Courier New",
        "myeongjo": "Batang",
        "pretendard": "Malgun Gothic",
        "cambriamath": "Cambria Math",
    }
    if key in id_map:
        return id_map[key]
    families = [part.strip().strip("'\"") for part in raw.split(",") if part.strip()]
    preferred = {
        "timesnewroman", "arial", "couriernew", "cambriamath",
        "malgungothic", "batang", "gulim", "calibri", "cambria",
        "liberationserif", "liberationsans", "liberationmono",
    }
    for family in families:
        fkey = re.sub(r"[^a-z0-9]", "", family.lower())
        if fkey in preferred:
            return family
    for family in families:
        fkey = re.sub(r"[^a-z0-9]", "", family.lower())
        if family and not family.lower().startswith("sdy ") and fkey not in ("serif", "sansserif", "monospace"):
            return family
    if "mono" in key or "courier" in key:
        return "Courier New"
    if "sans" in key or "helv" in key:
        return "Arial"
    if any("\uac00" <= c <= "\ud7af" for c in raw):
        return "Malgun Gothic"
    return "Times New Roman"


def _css_number(value, default=None):
    try:
        return float(re.sub(r"[^0-9.+-]", "", str(value)))
    except (TypeError, ValueError):
        return default


class _TightSpanParser(HTMLParser):
    """Extract the importer's absolutely positioned PDF spans.

    SDYnotes displays PDF text as positioned spans inside one text element.  A
    single flowing Word textbox ignores those ``left/top/data-pdf-w`` values, so
    the DOCX converter must keep them and create small anchored textboxes.
    """
    def __init__(self, default_font="Arial", default_px=12):
        super().__init__(convert_charrefs=True)
        self.default_font = _word_font(default_font or "Arial")
        self.default_px = float(default_px or 12)
        self.spans = []
        self._active = None
        self._depth = 0

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        attrs_d = dict(attrs)
        if tag == "span" and self._active is None:
            styles = _HtmlRuns._style(attrs)
            px = _css_number(attrs_d.get("data-fs") or styles.get("font-size"), self.default_px)
            font = _word_font(styles.get("font-family") or self.default_font)
            color = None
            match = re.search(r"#([0-9a-f]{6})", styles.get("color", ""), re.I)
            if match:
                color = match.group(1)
            self._active = {
                "text": "",
                "left": _css_number(styles.get("left"), 0.0),
                "top": _css_number(styles.get("top"), None),
                "width": _css_number(attrs_d.get("data-pdf-w"), None),
                "base": _css_number(attrs_d.get("data-pdf-base"), None),
                "style": {
                    "font": font,
                    "px": px,
                    "bold": styles.get("font-weight", "").lower() in ("bold", "700", "800", "900"),
                    "italic": styles.get("font-style", "").lower() == "italic",
                    "underline": False,
                    "color": color,
                },
            }
            self._depth = 1
            return
        if self._active is not None:
            self._depth += 1
            if tag in ("b", "strong"):
                self._active["style"]["bold"] = True
            elif tag in ("i", "em") and "zsp" not in (attrs_d.get("class") or "").split():
                self._active["style"]["italic"] = True
            elif tag == "u":
                self._active["style"]["underline"] = True

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self._active is None:
            return
        self._depth -= 1
        if self._depth <= 0:
            active = self._active
            self._active = None
            text = active.get("text") or ""
            if text:
                px = float(active["style"].get("px") or self.default_px)
                # CSS top is the browser positioning source.  If an older saved
                # element lacks it, reconstruct it from baseline like pdf_layout.
                top = active.get("top")
                if top is None and active.get("base") is not None:
                    top = float(active["base"]) - px * .8
                active["top"] = 0.0 if top is None else float(top)
                self.spans.append(active)

    def handle_data(self, data):
        if self._active is not None and data:
            self._active["text"] += data

    def handle_entityref(self, name):
        self.handle_data(html.unescape(f"&{name};"))


def _tight_spans(value, element):
    parser = _TightSpanParser(element.get("font") or "Arial", element.get("fontSize") or 12)
    try:
        parser.feed(value or "")
        parser.close()
    except Exception:
        return []
    return parser.spans

def _html_runs(value, element):
    parser = _HtmlRuns(element.get("font") or "Arial", element.get("fontSize") or 12)
    try:
        parser.feed(value or "")
        parser.close()
    except Exception:
        parser.runs = [(re.sub(r"<[^>]+>", "", value or ""), parser.default)]
    return parser.runs


def _latex_plain(value):
    """Readable Word fallback for the parser's already-correct LaTeX element."""
    text = str(value or "")
    greek = {
        "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
        "theta": "θ", "lambda": "λ", "mu": "μ", "pi": "π", "sigma": "σ",
        "phi": "φ", "psi": "ψ", "omega": "ω", "Gamma": "Γ", "Delta": "Δ",
        "Lambda": "Λ", "Pi": "Π", "Sigma": "Σ", "Phi": "Φ", "Psi": "Ψ",
        "Omega": "Ω",
    }
    symbols = {
        "infty": "∞", "partial": "∂", "nabla": "∇", "int": "∫", "oint": "∮",
        "sum": "∑", "prod": "∏", "sqrt": "√", "approx": "≈", "ne": "≠",
        "le": "≤", "ge": "≥", "pm": "±", "times": "×", "cdot": "·", "in": "∈",
        "notin": "∉", "to": "→", "leftarrow": "←", "Rightarrow": "⇒",
        "Leftrightarrow": "⇔", "forall": "∀", "exists": "∃", "ell": "ℓ",
        "hbar": "ℏ", "mathbb": "", "mathrm": "", "mathbf": "", "text": "",
    }

    # Make the common structural operators readable before removing command
    # names. Repeating the innermost balanced replacement also handles nested
    # fractions without introducing a second LaTeX parser.
    for _ in range(8):
        match = re.search(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}", text)
        if match:
            text = text[:match.start()] + f"({match.group(1)})⁄({match.group(2)})" + text[match.end():]
            continue
        match = re.search(r"\\sqrt\s*\{([^{}]*)\}", text)
        if match:
            text = text[:match.start()] + f"√({match.group(1)})" + text[match.end():]
            continue
        break

    for name, symbol in {**greek, **symbols}.items():
        text = re.sub(r"\\" + re.escape(name) + r"\b", symbol, text)
    text = re.sub(r"\\(?:left|right|displaystyle|textstyle|quad|,|;|!| |operatorname)\b", "", text)
    text = re.sub(r"\\text\s*\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\(?:mathbb|mathrm|mathbf|mathcal|mathsf|mathtt|mathit)\s*\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\^\{([^{}]*)\}", r"^(\1)", text)
    text = re.sub(r"_\{([^{}]*)\}", r"_(\1)", text)
    text = text.replace("\\", "").replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text or "Equation"



_MATH_SYMBOLS = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
    "varepsilon": "ε", "zeta": "ζ", "eta": "η", "theta": "θ", "vartheta": "ϑ",
    "iota": "ι", "kappa": "κ", "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ",
    "pi": "π", "rho": "ρ", "varrho": "ϱ", "sigma": "σ", "tau": "τ", "upsilon": "υ",
    "phi": "φ", "varphi": "φ", "chi": "χ", "psi": "ψ", "omega": "ω",
    "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ", "Pi": "Π",
    "Sigma": "Σ", "Upsilon": "Υ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω",
    "infty": "∞", "partial": "∂", "nabla": "∇", "approx": "≈", "sim": "∼",
    "simeq": "≃", "cong": "≅", "ne": "≠", "neq": "≠", "le": "≤", "leq": "≤",
    "ge": "≥", "geq": "≥", "pm": "±", "mp": "∓", "times": "×", "cdot": "·",
    "div": "÷", "circ": "∘", "bullet": "∙", "in": "∈", "notin": "∉", "subset": "⊂",
    "subseteq": "⊆", "supset": "⊃", "supseteq": "⊇", "cup": "∪", "cap": "∩",
    "forall": "∀", "exists": "∃", "neg": "¬", "land": "∧", "lor": "∨", "to": "→",
    "rightarrow": "→", "leftarrow": "←", "Rightarrow": "⇒", "Leftarrow": "⇐",
    "Leftrightarrow": "⇔", "mapsto": "↦", "ell": "ℓ", "hbar": "ℏ", "Re": "ℜ", "Im": "ℑ",
    "int": "∫", "oint": "∮", "sum": "∑", "prod": "∏", "lim": "lim", "min": "min", "max": "max",
}

_MATH_SKIP_COMMANDS = {
    "left", "right", "displaystyle", "textstyle", "scriptstyle", "scriptscriptstyle",
    "quad", "qquad", "!", ",", ";", ":", " ", "mathrm", "mathbf", "mathit",
    "mathsf", "mathtt", "mathbb", "mathcal", "operatorname", "text",
}


def _m_el(tag, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(key) if ":" in key else key, str(value))
    return el


def _m_run(text):
    from docx.oxml.ns import qn
    r = _m_el("m:r")
    t = _m_el("m:t")
    value = str(text or "")
    if value[:1].isspace() or value[-1:].isspace() or "  " in value:
        t.set(qn("xml:space"), "preserve")
    t.text = value
    r.append(t)
    return r


def _m_arg(tag, children):
    node = _m_el(tag)
    for child in children or [_m_run("")]:
        node.append(child)
    return node


class _LatexOmmlParser:
    """Small LaTeX-to-OMML converter for imported PDF equations.

    This is intentionally conservative: it covers the structures the importer
    emits most often (fractions, roots, scripts, Greek/symbol commands and text
    wrappers).  Unknown commands remain readable inside a real Word equation
    object instead of being flattened to an ordinary textbox.
    """
    def __init__(self, source):
        self.s = self._strip_math(str(source or ""))
        self.i = 0

    @staticmethod
    def _strip_math(value):
        value = value.strip()
        value = re.sub(r"^\\\((.*)\\\)$", r"\1", value, flags=re.S)
        value = re.sub(r"^\\\[(.*)\\\]$", r"\1", value, flags=re.S)
        if value.startswith("$$") and value.endswith("$$") and len(value) >= 4:
            value = value[2:-2]
        elif value.startswith("$") and value.endswith("$") and len(value) >= 2:
            value = value[1:-1]
        return value.strip()

    def _peek(self):
        return self.s[self.i] if self.i < len(self.s) else ""

    def _skip_ws(self):
        while self.i < len(self.s) and self.s[self.i].isspace():
            self.i += 1

    def parse(self, stop=""):
        out = []
        buf = []
        while self.i < len(self.s):
            ch = self._peek()
            if stop and ch == stop:
                break
            if ch == "\\" or ch in "{}^_":
                if ch in "^_":
                    # Attach scripts to the preceding atom.  In a common case
                    # like ``x_i^2`` the base may still be the last buffered
                    # character, not an emitted OMML node yet.
                    if buf:
                        if len(buf) > 1:
                            out.append(_m_run("".join(buf[:-1])))
                        base = [_m_run(buf[-1])]
                        buf = []
                    elif out:
                        base = [out.pop()]
                    else:
                        base = [_m_run("")]
                    out.extend(self._with_scripts(base))
                    continue
                if buf:
                    out.append(_m_run("".join(buf)))
                    buf = []
                if ch == "{":
                    self.i += 1
                    out.extend(self.parse("}"))
                    if self._peek() == "}":
                        self.i += 1
                    continue
                if ch == "}":
                    break
                atom = self._atom()
                out.extend(self._with_scripts(atom))
            else:
                buf.append(ch)
                self.i += 1
        if buf:
            out.append(_m_run("".join(buf)))
        return out

    def _group(self):
        self._skip_ws()
        if self._peek() == "{":
            self.i += 1
            children = self.parse("}")
            if self._peek() == "}":
                self.i += 1
            return children
        return self._atom()

    def _command(self):
        if self._peek() != "\\":
            return ""
        self.i += 1
        if self.i >= len(self.s):
            return ""
        if not self.s[self.i].isalpha():
            ch = self.s[self.i]
            self.i += 1
            return ch
        start = self.i
        while self.i < len(self.s) and self.s[self.i].isalpha():
            self.i += 1
        return self.s[start:self.i]

    def _atom(self):
        self._skip_ws()
        ch = self._peek()
        if not ch:
            return []
        if ch == "{":
            self.i += 1
            children = self.parse("}")
            if self._peek() == "}":
                self.i += 1
            return children
        if ch == "\\":
            name = self._command()
            if name == "frac":
                f = _m_el("m:f")
                f.append(_m_arg("m:num", self._group()))
                f.append(_m_arg("m:den", self._group()))
                return [f]
            if name == "sqrt":
                rad = _m_el("m:rad")
                rpr = _m_el("m:radPr")
                deg_hide = _m_el("m:degHide", **{"m:val": "1"})
                rpr.append(deg_hide)
                rad.append(rpr)
                rad.append(_m_arg("m:e", self._group()))
                return [rad]
            if name in ("overline", "bar"):
                bar = _m_el("m:bar")
                bpr = _m_el("m:barPr")
                pos = _m_el("m:pos", **{"m:val": "top"})
                bpr.append(pos); bar.append(bpr); bar.append(_m_arg("m:e", self._group()))
                return [bar]
            if name in ("underline",):
                bar = _m_el("m:bar")
                bpr = _m_el("m:barPr")
                pos = _m_el("m:pos", **{"m:val": "bot"})
                bpr.append(pos); bar.append(bpr); bar.append(_m_arg("m:e", self._group()))
                return [bar]
            if name in ("mathrm", "mathbf", "mathit", "mathsf", "mathtt", "mathbb", "mathcal", "text", "operatorname"):
                return self._group()
            if name in ("left", "right"):
                self._skip_ws()
                if self._peek() == ".":
                    self.i += 1
                    return []
                if self._peek():
                    ch2 = self._peek(); self.i += 1
                    return [_m_run(ch2)]
                return []
            if name in _MATH_SYMBOLS:
                return [_m_run(_MATH_SYMBOLS[name])]
            if name in _MATH_SKIP_COMMANDS:
                return []
            return [_m_run(name if len(name) > 1 else name)]
        self.i += 1
        if ch in "^_}":
            return [_m_run(ch)]
        return [_m_run(ch)]

    def _script_arg(self):
        self._skip_ws()
        if self._peek() == "{":
            self.i += 1
            children = self.parse("}")
            if self._peek() == "}":
                self.i += 1
            return children
        return self._atom()

    def _with_scripts(self, base):
        sup = sub = None
        for _ in range(2):
            self._skip_ws()
            ch = self._peek()
            if ch not in "^_":
                break
            self.i += 1
            if ch == "^":
                sup = self._script_arg()
            else:
                sub = self._script_arg()
        if sup is None and sub is None:
            return base
        if sup is not None and sub is not None:
            node = _m_el("m:sSubSup")
            node.append(_m_arg("m:e", base))
            node.append(_m_arg("m:sub", sub))
            node.append(_m_arg("m:sup", sup))
            return [node]
        if sup is not None:
            node = _m_el("m:sSup")
            node.append(_m_arg("m:e", base))
            node.append(_m_arg("m:sup", sup))
            return [node]
        node = _m_el("m:sSub")
        node.append(_m_arg("m:e", base))
        node.append(_m_arg("m:sub", sub))
        return [node]


def _latex_to_omml_children(latex):
    children = _LatexOmmlParser(latex).parse()
    return children or [_m_run(_latex_plain(latex))]

def _emu_px(value):
    # The parser canvas is 96 CSS px/in.  Word drawing coordinates are EMU.
    return max(0, int(round(float(value) * 914400 / 96)))


def _add_text_box(paragraph, element, z, counter):
    """Add an editable, absolutely positioned VML text box to a page."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree

    x, y = float(element.get("x") or 0), float(element.get("y") or 0)
    w, h = max(2.0, float(element.get("w") or 2)), max(2.0, float(element.get("h") or 2))
    runs = _html_runs(element.get("html") or "", element)
    if not runs:
        return
    run = paragraph.add_run()
    pict = OxmlElement("w:pict")
    vml = "{urn:schemas-microsoft-com:vml}"
    shape = etree.Element(vml + "shape")
    shape.set("id", f"sdyText{counter}")
    shape.set("type", "#_x0000_t202")
    shape.set("stroked", "f")
    shape.set("filled", "f")
    shape.set("style", (
        f"position:absolute;margin-left:{x * .75:.3f}pt;margin-top:{y * .75:.3f}pt;"
        f"width:{w * .75:.3f}pt;height:{h * .75:.3f}pt;"
        "mso-position-horizontal:absolute;mso-position-horizontal-relative:page;"
        "mso-position-vertical:absolute;mso-position-vertical-relative:page;"
        "mso-wrap-style:none;z-index:%d" % z
    ))
    textbox = etree.SubElement(shape, vml + "textbox")
    textbox.set("inset", "0,0,0,0")
    content = OxmlElement("w:txbxContent")
    inner = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    line_px = max([float(style.get("px") or element.get("fontSize") or 12) for _, style in runs] or [float(element.get("fontSize") or 12)])
    # Word uses twentieths of a point.  Match the importer/browser line box
    # instead of forcing every textbox to Word's default 12pt line spacing.
    spacing.set(qn("w:line"), str(max(20, int(round(line_px * .75 * 20 * 1.08)))))
    spacing.set(qn("w:lineRule"), "exact")
    ppr.append(spacing)
    inner.append(ppr)
    for text, style in runs:
        wr = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        font = str(style.get("font") or "Arial")
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), font)
        rpr.append(fonts)
        size_pt = max(1.0, float(style.get("px") or 12) * .75)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(max(2, int(round(size_pt * 2)))))
        rpr.append(sz)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(max(2, int(round(size_pt * 2)))))
        rpr.append(szcs)
        if style.get("bold"):
            rpr.append(OxmlElement("w:b"))
        if style.get("italic"):
            rpr.append(OxmlElement("w:i"))
        if style.get("underline"):
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rpr.append(u)
        if style.get("color"):
            color = OxmlElement("w:color")
            color.set(qn("w:val"), str(style["color"]).replace("#", ""))
            rpr.append(color)
        wr.append(rpr)
        pieces = str(text).split("\n")
        for i, piece in enumerate(pieces):
            if i:
                inner.append(wr)
                wr = OxmlElement("w:r")
                wr.append(deepcopy(rpr))
                br = OxmlElement("w:br")
                wr.append(br)
            wt = OxmlElement("w:t")
            if piece[:1].isspace() or piece[-1:].isspace() or "  " in piece:
                wt.set(qn("xml:space"), "preserve")
            wt.text = piece
            wr.append(wt)
        inner.append(wr)
    content.append(inner)
    textbox.append(content)
    shape.append(textbox)
    pict.append(shape)
    run._r.append(pict)




def _add_math_box(paragraph, element, z, counter):
    """Add an absolutely positioned editable Word equation (OMML)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree

    x, y = float(element.get("x") or 0), float(element.get("y") or 0)
    w, h = max(2.0, float(element.get("w") or 2)), max(2.0, float(element.get("h") or 2))
    latex = str(element.get("latex") or "").strip()
    if not latex:
        return False
    try:
        math_children = _latex_to_omml_children(latex)
    except Exception:
        math_children = [_m_run(_latex_plain(latex))]
    run = paragraph.add_run()
    pict = OxmlElement("w:pict")
    vml = "{urn:schemas-microsoft-com:vml}"
    shape = etree.Element(vml + "shape")
    shape.set("id", f"sdyMath{counter}")
    shape.set("type", "#_x0000_t202")
    shape.set("stroked", "f")
    shape.set("filled", "f")
    shape.set("style", (
        f"position:absolute;margin-left:{x * .75:.3f}pt;margin-top:{y * .75:.3f}pt;"
        f"width:{w * .75:.3f}pt;height:{h * .75:.3f}pt;"
        "mso-position-horizontal:absolute;mso-position-horizontal-relative:page;"
        "mso-position-vertical:absolute;mso-position-vertical-relative:page;"
        "mso-wrap-style:none;z-index:%d" % z
    ))
    textbox = etree.SubElement(shape, vml + "textbox")
    textbox.set("inset", "0,0,0,0")
    content = OxmlElement("w:txbxContent")
    inner = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), str(max(20, int(round(float(element.get("fontSize") or 12) * .75 * 20 * 1.15)))))
    spacing.set(qn("w:lineRule"), "exact")
    ppr.append(spacing)
    inner.append(ppr)
    omath = OxmlElement("m:oMath")
    for child in math_children:
        omath.append(child)
    if element.get("displayMath"):
        para = OxmlElement("m:oMathPara")
        para.append(omath)
        inner.append(para)
    else:
        inner.append(omath)
    content.append(inner)
    textbox.append(content)
    shape.append(textbox)
    pict.append(shape)
    run._r.append(pict)
    return True

def _tight_span_html(text, style):
    css = [
        f"font-family:{style.get('font') or 'Arial'}",
        f"font-size:{float(style.get('px') or 12):.3f}px",
        f"font-weight:{700 if style.get('bold') else 400}",
        f"font-style:{'italic' if style.get('italic') else 'normal'}",
    ]
    if style.get("color"):
        css.append(f"color:#{str(style.get('color')).replace('#', '')}")
    body = _escape_text(text)
    if style.get("underline"):
        body = f"<u>{body}</u>"
    return f"<span style=\"{html.escape(';'.join(css), quote=True)}\">{body}</span>"


def _add_tight_text_boxes(paragraph, element, z, counter):
    """Preserve SDYnotes PDF text geometry when writing DOCX.

    Imported PDF text in the site is not a normal flowing paragraph: every word
    span has absolute CSS coordinates plus the original PDF advance width.  If
    we put the whole element into one Word textbox, Word reflows it with its own
    line-breaking/font fallback and the page no longer matches the site.  Emit
    small page-anchored textboxes for those spans instead.
    """
    spans = _tight_spans(element.get("html") or "", element)
    if not spans:
        return 0
    base_x = float(element.get("x") or 0)
    base_y = float(element.get("y") or 0)
    used = 0
    for sp in spans:
        text = re.sub(r"\s+$", "", str(sp.get("text") or ""))
        if not text:
            continue
        style = sp.get("style") or {}
        px = max(1.0, float(style.get("px") or element.get("fontSize") or 12))
        width = sp.get("width")
        if not width or width <= 0:
            width = max(2.0, len(text) * px * .55)
        # A couple of pixels of slack avoids Word clipping/substituting a final
        # glyph while the absolute x positions still carry inter-word spacing.
        width = max(2.0, float(width) + 2.0)
        piece = {
            "type": "text",
            "x": base_x + float(sp.get("left") or 0),
            "y": base_y + float(sp.get("top") or 0),
            "w": width,
            "h": max(2.0, px * 1.25),
            "font": style.get("font") or element.get("font") or "Arial",
            "fontSize": px,
            "html": _tight_span_html(text, style),
        }
        _add_text_box(paragraph, piece, z, counter + used)
        used += 1
    return used

def _add_floating_picture(paragraph, path, element, z, counter):
    """Insert a normal DOCX image and move its inline drawing to page coordinates."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    x, y = float(element.get("x") or 0), float(element.get("y") or 0)
    w, h = max(1.0, float(element.get("w") or 1)), max(1.0, float(element.get("h") or 1))
    run = paragraph.add_run()
    try:
        run.add_picture(path, width=Inches(w / 96), height=Inches(h / 96))
    except Exception:
        return
    drawing = run._r.drawing_lst[-1]
    inline = drawing.find(qn("wp:inline"))
    if inline is None:
        return
    children = list(inline)
    anchor = OxmlElement("wp:anchor")
    for key, value in (("distT", "0"), ("distB", "0"), ("distL", "0"),
                       ("distR", "0"), ("simplePos", "0"),
                       ("relativeHeight", str(max(0, int(z)))),
                       ("behindDoc", "1" if z == 0 else "0"),
                       ("locked", "0"), ("layoutInCell", "1"),
                       ("allowOverlap", "1")):
        anchor.set(qn(f"wp:{key}"), value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0"); simple.set("y", "0")
    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "page")
    off_h = OxmlElement("wp:posOffset")
    off_h.text = str(_emu_px(x))
    pos_h.append(off_h)
    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "page")
    off_v = OxmlElement("wp:posOffset")
    off_v.text = str(_emu_px(y))
    pos_v.append(off_v)
    extent = OxmlElement("wp:extent")
    extent.set("cx", str(_emu_px(w)))
    extent.set("cy", str(_emu_px(h)))
    effect = OxmlElement("wp:effectExtent")
    for side in ("l", "t", "r", "b"):
        effect.set(side, "0")
    wrap = OxmlElement("wp:wrapNone")
    anchor.extend([simple, pos_h, pos_v, extent, effect, wrap])
    for child in children:
        if child.tag == qn("wp:extent"):
            continue
        if child.tag == qn("wp:docPr"):
            child.set("id", str(counter))
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def _set_section(section, width_px, height_px):
    from docx.shared import Inches
    zero = Inches(0)
    section.page_width = Inches(width_px / 96)
    section.page_height = Inches(height_px / 96)
    section.left_margin = zero
    section.right_margin = zero
    section.top_margin = zero
    section.bottom_margin = zero
    section.header_distance = zero
    section.footer_distance = zero


def _new_page_paragraph(document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    paragraph.paragraph_format.line_spacing = 1
    return paragraph


def _build_docx(pages, target_w, target_h, title, jid, generated_assets):
    """Lay parser elements back over the parser's text-free page backgrounds."""
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.core_properties.title = title
    document.core_properties.subject = "SDYnotes PDF parser export"
    first = True
    counter = 10
    for page in pages:
        if first:
            section = document.sections[0]
            first = False
        else:
            section = document.add_section(WD_SECTION_START.NEW_PAGE)
        _set_section(section, target_w, target_h)
        paragraph = _new_page_paragraph(document)
        elements = list(page.get("els") or [])
        # Background is the bottom layer. Figures are above it; parsed text and
        # equations are last so the text removed by the shared paint plan is not
        # reintroduced as a raster duplicate.
        elements.sort(key=lambda el: (
            0 if el.get("type") == "image" and el.get("isBg") else
            1 if el.get("type") == "image" else
            2 if el.get("type") == "latex" else 3
        ))
        for element in elements:
            kind = element.get("type")
            if kind == "image":
                path = _asset_path(element.get("url"), jid)
                if path:
                    generated_assets.add(path)
                    _add_floating_picture(paragraph, path, element,
                                          0 if element.get("isBg") else 20, counter)
                    counter += 1
            elif kind == "text":
                if element.get("tight") or element.get("pdfText"):
                    added = _add_tight_text_boxes(paragraph, element, 40, counter)
                    if added:
                        counter += added
                    else:
                        _add_text_box(paragraph, element, 40, counter)
                        counter += 1
                else:
                    _add_text_box(paragraph, element, 40, counter)
                    counter += 1
            elif kind == "latex":
                # Preserve imported/display/inline formulas as editable Word
                # equation objects (OMML), not as ordinary text or screenshots.
                if not _add_math_box(paragraph, element, 50, counter):
                    formula = dict(element)
                    formula["html"] = _escape_text(_latex_plain(element.get("latex")))
                    formula["font"] = "Cambria Math"
                    formula["fontSize"] = element.get("fontSize") or 12
                    _add_text_box(paragraph, formula, 50, counter)
                counter += 1

    # Remove the empty initial paragraph if possible only after content exists;
    # leaving it is harmless and avoids a malformed section boundary in older
    # Word readers.
    out = _output_path(jid)
    document.save(out)
    return out


def _escape_text(value):
    return html.escape(str(value or ""))


def _converter_worker(jid, src, original_name):
    generated_assets = set()
    try:
        # _pdf_one_page writes the same raster/vector assets as normal import.
        # Remember the directory before parsing so only this request's files
        # are cleaned up after the DOCX is complete.
        image_before = set(os.listdir(IMG_DIR))
        size_preset = _engine._pdf_size_preset(src)
        target_w, target_h = _engine._preset_dims(size_preset)
        _write_job(jid, status="working", phase="parsing", page=0,
                   total=0, name=original_name)
        # Use the same worker semaphore/accounting as the normal import worker;
        # only the output formatter is different.
        _engine._imp_sem.acquire()
        with _engine._imp_active_lock:
            _engine._imp_active += 1
        try:
            pages = _engine._imp_convert_pdf(src, jid, target_w, target_h)
        finally:
            with _engine._imp_active_lock:
                _engine._imp_active = max(0, _engine._imp_active - 1)
            _engine._imp_sem.release()
        _write_job(jid, status="working", phase="formatting",
                   page=len(pages), total=len(pages))
        # The parser writes image assets to the normal import image directory so
        # its established paint/image code is reused.  They are temporary for a
        # converter request and are removed after the DOCX has been assembled.
        for page in pages:
            for element in page.get("els", []):
                name = _asset_name(element.get("url"))
                if name and name not in image_before:
                    generated_assets.add(os.path.join(IMG_DIR, name))
        title = os.path.splitext(os.path.basename(original_name or "document.pdf"))[0][:100]
        out = _build_docx(pages, target_w, target_h, title, jid, generated_assets)
        _write_job(jid, status="done", phase="ready", page=len(pages),
                   total=len(pages), name=original_name, title=title,
                   pages=len(pages), bytes=os.path.getsize(out),
                   download=f"/api/converter/download/{jid}")
    except MemoryError:
        _write_job(jid, status="error", error="서버 메모리가 부족합니다. PDF를 나누어 다시 시도해 주세요")
    except Exception as exc:
        import traceback
        traceback.print_exc()
        _write_job(jid, status="error", error=f"변환 실패: {str(exc) or exc.__class__.__name__}")
    finally:
        try:
            if src and os.path.exists(src):
                os.remove(src)
        except OSError:
            pass
        # Asset files are not part of the converter's public state.  Use the
        # exact names emitted in this request, never an mtime-wide directory wipe.
        for path in generated_assets:
            try:
                if os.path.isfile(path):
                    os.remove(path)
            except OSError:
                pass
        # _imp_convert_pdf uses the import engine's progress helper internally;
        # its short-lived status file is not part of the standalone service.
        try:
            os.remove(_engine._imp_job_path(jid))
        except OSError:
            pass
        _sweep_converter_jobs()


@app.route("/api/converter/convert", methods=["POST"])
def converter_convert():
    _sweep_converter_jobs()
    upload = request.files.get("file")
    if not upload:
        return jsonify({"ok": False, "error": "PDF 파일을 선택해 주세요"}), 400
    jid = uuid.uuid4().hex[:16]
    src = os.path.join(CONVERTER_UPLOAD_DIR, f"{jid}.pdf")
    try:
        upload.save(src)
        size = os.path.getsize(src)
        if size <= 0:
            raise ValueError("빈 파일입니다")
        if size > CONVERTER_MAX_MB * 1024 * 1024:
            raise ValueError(f"PDF는 {CONVERTER_MAX_MB}MB 이하만 변환할 수 있습니다")
        with open(src, "rb") as fp:
            if fp.read(4) != b"%PDF":
                raise ValueError("PDF 파일만 올릴 수 있습니다")
        import pymupdf
        with pymupdf.open(src) as pdf:
            if pdf.is_encrypted and not pdf.authenticate(""):
                raise ValueError("비밀번호로 잠긴 PDF입니다. 잠금을 푼 뒤 다시 올려 주세요")
            count = pdf.page_count
            if count <= 0:
                raise ValueError("페이지가 없는 PDF입니다")
            if count > _engine.IMPORT_MAX_PAGES:
                raise ValueError(f"PDF는 {_engine.IMPORT_MAX_PAGES}쪽까지 변환할 수 있습니다")
    except Exception as exc:
        try:
            os.remove(src)
        except OSError:
            pass
        return jsonify({"ok": False, "error": str(exc)}), 400

    _write_job(jid, status="working", phase="queued", page=0, total=count,
               name=upload.filename or "document.pdf")
    threading.Thread(target=_converter_worker,
                     args=(jid, src, upload.filename or "document.pdf"),
                     daemon=True).start()
    return jsonify({"ok": True, "job": jid, "total": count})


@app.route("/api/converter/status", methods=["GET"])
def converter_status():
    jid = _safe_id(request.args.get("id"))
    if not jid:
        return jsonify({"ok": False, "error": "작업을 찾을 수 없습니다"}), 404
    job = _read_job(jid)
    if not job:
        return jsonify({"ok": False, "status": "gone", "error": "작업이 만료되었습니다"}), 404
    if job.get("status") == "done" and not os.path.isfile(_output_path(jid)):
        return jsonify({"ok": False, "status": "gone", "error": "변환 파일이 만료되었습니다"}), 404
    return jsonify({"ok": True, **{k: job.get(k) for k in (
        "status", "phase", "page", "total", "name", "title", "pages", "bytes", "download", "error"
    ) if k in job}})


@app.route("/api/converter/download/<jid>", methods=["GET"])
def converter_download(jid):
    jid = _safe_id(jid)
    path = _output_path(jid) if jid else ""
    if not path or not os.path.isfile(path):
        return jsonify({"ok": False, "error": "변환 파일이 없거나 만료되었습니다"}), 404
    job = _read_job(jid) or {}
    base = re.sub(r"[^0-9A-Za-z가-힣._ -]", "", str(job.get("title") or "converted"))[:80] or "converted"
    return send_file(path, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     as_attachment=True, download_name=f"{base}.docx",
                     max_age=0)
